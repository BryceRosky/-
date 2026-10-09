from __future__ import annotations

import re
import zipfile
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt


BUILD_DIR = Path(__file__).resolve().parent
DELIVERABLES_DIR = BUILD_DIR.parent
TEMPLATE = BUILD_DIR / "template" / "论文开题报告-Thesis+Proposal.docx"
CONTENT = DELIVERABLES_DIR / "thesis_proposal_content.md"
OUTPUT = DELIVERABLES_DIR / "thesis_proposal_luosiqi.docx"
RENDER_OUTPUT = DELIVERABLES_DIR / "thesis_proposal_luosiqi_render_compatible.docx"
FIGURE_UNIFIED_VALUE = DELIVERABLES_DIR / "assets" / "unified_value_framework.png"

TITLE = "面向离散扩散多模态大模型的生成优化方法研究"


def section_text(source: str, chinese_no: str, next_no: str | None) -> str:
    start = re.search(rf"^## {chinese_no}、.*$", source, re.M)
    if not start:
        raise ValueError(f"section {chinese_no} not found")
    end = re.search(rf"^## {next_no}、.*$", source[start.end() :], re.M) if next_no else None
    end_pos = start.end() + end.start() if end else len(source)
    return source[start.end() : end_pos].strip()


def parse_blocks(text: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if (
            line.startswith("|")
            and index + 1 < len(lines)
            and re.match(r"^\|?\s*:?-{3,}", lines[index + 1].strip())
        ):
            table_lines = [line, lines[index + 1].strip()]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            blocks.append(("table", "\n".join(table_lines)))
            continue
        if line.startswith("#### "):
            blocks.append(("heading4", line[5:].strip()))
        elif line.startswith("### "):
            blocks.append(("heading3", line[4:].strip()))
        elif line == "[[FIGURE:unified_value_framework]]":
            blocks.append(("figure", line))
        elif line.startswith("[[EQUATION:") and line.endswith("]]"):
            blocks.append(("equation", line[11:-2]))
        elif re.match(r"^[图表]\d+\s", line):
            blocks.append(("caption", line))
        elif line.startswith("## ") or line.startswith("# "):
            index += 1
            continue
        else:
            blocks.append(("body", line))
        index += 1
    return blocks


def remove_content(paragraph: Paragraph, preserve_permissions: bool = True) -> None:
    for child in list(paragraph._p):
        local = child.tag.split("}")[-1]
        if local == "pPr":
            continue
        if preserve_permissions and local in {"permStart", "permEnd"}:
            continue
        paragraph._p.remove(child)


def set_run_font(
    run,
    *,
    size: float = 12,
    bold: bool = False,
    east_asia: str = "宋体",
) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), "Times New Roman")


def format_paragraph(paragraph: Paragraph, kind: str) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.widow_control = True
    # A cloned paragraph may inherit keepNext from a heading. Reset it for
    # every block first, then enable it only for actual internal headings.
    pf.keep_with_next = False
    if kind in {"heading3", "heading4"}:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Pt(0)
        pf.left_indent = Pt(0)
        pf.line_spacing = 1.35
        pf.space_before = Pt(8 if kind == "heading3" else 5)
        pf.space_after = Pt(3)
        pf.keep_with_next = True
    elif kind == "reference":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.left_indent = Pt(24)
        pf.first_line_indent = Pt(-24)
        pf.line_spacing = 1.05
        pf.space_after = Pt(1)
    elif kind in {"figure", "caption", "table", "equation"}:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Pt(0)
        pf.left_indent = Pt(0)
        pf.line_spacing = 1.0
        pf.space_before = Pt(5 if kind == "figure" else 0)
        pf.space_after = Pt(2 if kind == "figure" else 6)
        pf.keep_with_next = kind in {"figure", "equation", "caption"}
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Pt(24)
        pf.left_indent = Pt(0)
        pf.line_spacing = 1.5


def style_run(run, kind: str) -> None:
    if kind == "heading3":
        set_run_font(run, size=12, bold=True, east_asia="黑体")
    elif kind == "heading4":
        set_run_font(run, size=11.5, bold=True, east_asia="黑体")
    elif kind == "reference":
        set_run_font(run, size=9.5, bold=False, east_asia="宋体")
    elif kind == "caption":
        set_run_font(run, size=10.5, bold=False, east_asia="宋体")
    else:
        set_run_font(run, size=12, bold=False, east_asia="宋体")


INLINE_MATH = re.compile(
    r"([A-Za-z\u0370-\u03ff\u1f00-\u1fff\u0300-\u036f]+)"
    r"([_^])"
    r"([A-Za-z0-9\u0370-\u03ff\u1f00-\u1fff+\-]+(?:_[A-Za-z0-9\u0370-\u03ff\u1f00-\u1fff+\-]+)*)"
)


def add_inline_text(paragraph: Paragraph, text: str, kind: str) -> None:
    cursor = 0
    for match in INLINE_MATH.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            style_run(run, kind)
        base = paragraph.add_run(match.group(1))
        style_run(base, kind)
        index = paragraph.add_run(match.group(3).replace("_", ""))
        style_run(index, kind)
        if match.group(2) == "_":
            index.font.subscript = True
        else:
            index.font.superscript = True
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        style_run(run, kind)


def put_text(paragraph: Paragraph, text: str, kind: str) -> None:
    remove_content(paragraph, preserve_permissions=True)
    add_inline_text(paragraph, text, kind)
    format_paragraph(paragraph, kind)
    # Keep the contribution table and its caption together on the final page.
    if kind == "caption" and text.startswith("表1 "):
        paragraph.paragraph_format.page_break_before = True


def put_figure(paragraph: Paragraph) -> None:
    if not FIGURE_UNIFIED_VALUE.exists():
        raise FileNotFoundError(FIGURE_UNIFIED_VALUE)
    remove_content(paragraph, preserve_permissions=True)
    run = paragraph.add_run()
    run.add_picture(str(FIGURE_UNIFIED_VALUE), width=Inches(5.95))
    format_paragraph(paragraph, "figure")


def omath_run(text: str, *, normal: bool = False):
    run = OxmlElement("m:r")
    run_pr = OxmlElement("m:rPr")
    if normal:
        run_pr.append(OxmlElement("m:nor"))
    run.append(run_pr)
    text_el = OxmlElement("m:t")
    text_el.text = text
    run.append(text_el)
    return run


def omath_script(base: str, sub: str | None = None, sup: str | None = None):
    if sub is not None and sup is not None:
        node = OxmlElement("m:sSubSup")
        node.append(OxmlElement("m:sSubSupPr"))
    elif sub is not None:
        node = OxmlElement("m:sSub")
        node.append(OxmlElement("m:sSubPr"))
    elif sup is not None:
        node = OxmlElement("m:sSup")
        node.append(OxmlElement("m:sSupPr"))
    else:
        return omath_run(base)
    base_el = OxmlElement("m:e")
    base_el.append(omath_run(base))
    node.append(base_el)
    if sub is not None:
        sub_el = OxmlElement("m:sub")
        sub_el.append(omath_run(sub, normal=len(sub) > 1))
        node.append(sub_el)
    if sup is not None:
        sup_el = OxmlElement("m:sup")
        sup_el.append(omath_run(sup, normal=len(sup) > 1))
        node.append(sup_el)
    return node


def build_equation(key: str):
    math = OxmlElement("m:oMath")

    def add_text(value: str, *, normal: bool = False) -> None:
        math.append(omath_run(value, normal=normal))

    def add_script(base: str, sub: str | None = None, sup: str | None = None) -> None:
        math.append(omath_script(base, sub=sub, sup=sup))

    if key == "state_action":
        add_script("s", "t"); add_text(" = (")
        add_script("z", "t"); add_text(", ")
        add_script("M", "t"); add_text(", t, ")
        add_script("b", "t"); add_text(", ")
        add_script("h", "t"); add_text("),    ")
        add_script("a", "t"); add_text(" = (")
        add_script("S", "t"); add_text(", ")
        add_script("y", "t"); add_text("(")
        add_script("S", "t"); add_text(")),    x = ")
        add_text("Dec", normal=True); add_text("(")
        add_script("z", "T"); add_text(")")
    elif key == "conditional_value":
        add_script("V", "R,eval", "π,cont"); add_text("(")
        add_script("s", "t"); add_text(", c) = ")
        add_script("E", "π,cont"); add_text("[")
        add_script("R", "eval"); add_text("(x, c) | ")
        add_script("s", "t"); add_text(", c]")
    elif key == "constraints":
        add_text("E[C(𝒯)] ≤ "); add_script("B", "comp"); add_text(",    ")
        add_script("U", "k"); add_text(" ≥ ")
        add_script("U", "k"); add_text("(baseline) − ")
        add_script("ε", "k"); add_text(",    ")
        add_script("D", "div"); add_text(" ≥ ")
        add_script("D", "div"); add_text("(baseline) − ")
        add_script("δ", "div")
    else:
        raise ValueError(f"unknown equation: {key}")
    return math


def put_equation(paragraph: Paragraph, key: str) -> None:
    remove_content(paragraph, preserve_permissions=True)
    paragraph._p.append(build_equation(key))
    format_paragraph(paragraph, "equation")


def parse_table_rows(text: str) -> list[list[str]]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    rows: list[list[str]] = []
    for idx, line in enumerate(lines):
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if idx == 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    if len(rows) < 2 or len({len(row) for row in rows}) != 1:
        raise ValueError("invalid markdown table")
    return rows


def set_cell_margins(cell, value: int = 55) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def put_table(paragraph: Paragraph, text: str) -> None:
    rows = parse_table_rows(text)
    remove_content(paragraph, preserve_permissions=True)
    format_paragraph(paragraph, "table")
    table = paragraph._parent.add_table(
        rows=len(rows), cols=len(rows[0]), width=Inches(6.25)
    )
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    widths = [0.95, 1.75, 1.75, 1.8] if len(rows[0]) == 4 else [6.25 / len(rows[0])] * len(rows[0])
    for row_idx, values in enumerate(rows):
        tr_pr = table.rows[row_idx]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if row_idx == 0:
            header = OxmlElement("w:tblHeader")
            tr_pr.append(header)
        for col_idx, value in enumerate(values):
            cell = table.cell(row_idx, col_idx)
            cell.width = Inches(widths[col_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "D9EAF2")
                cell._tc.get_or_add_tcPr().append(shading)
            cell_p = cell.paragraphs[0]
            remove_content(cell_p, preserve_permissions=False)
            cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            cell_p.paragraph_format.space_before = Pt(0)
            cell_p.paragraph_format.space_after = Pt(0)
            cell_p.paragraph_format.line_spacing = 1.05
            cell_run = cell_p.add_run(value)
            set_run_font(
                cell_run,
                size=8.2 if row_idx else 8.5,
                bold=row_idx == 0,
                east_asia="黑体" if row_idx == 0 else "宋体",
            )
    paragraph._p.addnext(table._tbl)


def clone_blank(source: Paragraph, after: Paragraph) -> Paragraph:
    new_p = deepcopy(source._p)
    for child in list(new_p):
        if child.tag.split("}")[-1] != "pPr":
            new_p.remove(child)
    ppr = new_p.find(qn("w:pPr"))
    if ppr is not None:
        keep_next = ppr.find(qn("w:keepNext"))
        if keep_next is not None:
            ppr.remove(keep_next)
    # A generated table is inserted immediately after its placeholder
    # paragraph.  Continue after that table so later paragraphs do not slip
    # between the caption and the table.
    next_element = after._p.getnext()
    if next_element is not None and next_element.tag == qn("w:tbl"):
        next_element.addnext(new_p)
    else:
        after._p.addnext(new_p)
    return Paragraph(new_p, after._parent)


def remove_paragraph_numbering(paragraph: Paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is not None:
        ppr.remove(num_pr)


def rebuild_official_heading(paragraph: Paragraph) -> None:
    """Rebuild visible heading text while retaining numbering and pPr.

    The template splits some official headings across several runs. In
    particular, LibreOffice omits “预期成果” from heading 5 when exporting the
    original run structure. A single visible run avoids that compatibility
    problem without replacing the paragraph or its numId.
    """
    text = paragraph.text
    remove_content(paragraph, preserve_permissions=True)
    run = paragraph.add_run(text)
    set_run_font(run, size=12, bold=True, east_asia="仿宋_GB2312")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def fill_range(slots: list[Paragraph], blocks: list[tuple[str, str]]) -> None:
    if not slots:
        raise ValueError("no paragraph slots")
    clone_source = slots[-1]
    for idx, (kind, text) in enumerate(blocks):
        if idx < len(slots):
            paragraph = slots[idx]
        else:
            paragraph = clone_blank(clone_source, slots[-1])
            slots.append(paragraph)
        if kind == "figure":
            put_figure(paragraph)
        elif kind == "equation":
            put_equation(paragraph, text)
        elif kind == "table":
            put_table(paragraph, text)
        else:
            put_text(paragraph, text, kind)
    for extra in slots[len(blocks) :]:
        remove_content(extra, preserve_permissions=True)


def insert_run_before_perm_end(paragraph: Paragraph, text: str, size: float = 12) -> None:
    # Preserve the editable-range markers while replacing the visible cell value.
    for child in list(paragraph._p):
        local = child.tag.split("}")[-1]
        if local in {"r", "hyperlink", "fldSimple", "sdt"}:
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    set_run_font(run, size=size)
    perm_end = paragraph._p.find(qn("w:permEnd"))
    if perm_end is not None:
        paragraph._p.remove(run._r)
        perm_end.addprevious(run._r)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def set_dropdown(cell, value: str) -> None:
    sdts = cell._tc.xpath(".//w:sdt")
    if len(sdts) != 1:
        raise ValueError(f"expected one dropdown, found {len(sdts)}")
    sdt = sdts[0]
    sdt_pr = sdt.find(qn("w:sdtPr"))
    placeholder = sdt_pr.find(qn("w:showingPlcHdr")) if sdt_pr is not None else None
    if placeholder is not None:
        sdt_pr.remove(placeholder)
    texts = list(sdt.iter(qn("w:t")))
    if not texts:
        raise ValueError("dropdown has no display text")
    texts[0].text = value
    for extra in texts[1:]:
        extra.text = ""
    run = texts[0].getparent()
    rpr = run.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        run.insert(0, rpr)
    for tag in ("w:rStyle", "w:color"):
        old = rpr.find(qn(tag))
        if old is not None:
            rpr.remove(old)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    rpr.append(color)
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")
    fonts.set(qn("w:cs"), "Times New Roman")


def set_source_checkbox(cell, selected_index: int) -> None:
    sdts = cell._tc.xpath(".//w:sdt")
    if len(sdts) != 7:
        raise ValueError(f"expected seven source checkboxes, found {len(sdts)}")
    for idx, sdt in enumerate(sdts):
        checked = sdt.find(".//" + qn("w14:checked"))
        if checked is None:
            raise ValueError("checkbox lacks w14:checked")
        is_selected = idx == selected_index
        checked.set(qn("w14:val"), "1" if is_selected else "0")
        texts = list(sdt.iter(qn("w:t")))
        if texts:
            texts[0].text = "√" if is_selected else "☐"


def set_signature_date_placeholder(paragraph: Paragraph, value: str = "____________") -> None:
    start = None
    end = None
    for child in list(paragraph._p):
        if child.tag == qn("w:permStart") and child.get(qn("w:id")) == "895295510":
            start = child
        if child.tag == qn("w:permEnd") and child.get(qn("w:id")) == "895295510":
            end = child
    if start is None or end is None:
        raise ValueError("signature date editable range not found")
    cursor = start.getnext()
    while cursor is not None and cursor is not end:
        nxt = cursor.getnext()
        paragraph._p.remove(cursor)
        cursor = nxt
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "仿宋_GB2312")
    fonts.set(qn("w:cs"), "Times New Roman")
    rpr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    rpr.append(size)
    run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.text = value
    run.append(text_el)
    end.addprevious(run)


def make_render_safe_copy(source: Path, destination: Path) -> None:
    """Map requested Office fonts to installed fonts for deterministic PDF export."""
    replacements = [
        ("仿宋_GB2312", "Noto Sans CJK SC"),
        ("FangSong_GB2312", "Noto Sans CJK SC"),
        ("FangSong", "Noto Sans CJK SC"),
        ("黑体", "Noto Sans CJK SC"),
        ("SimHei", "Noto Sans CJK SC"),
        ("微软雅黑", "Noto Sans CJK SC"),
        ("Microsoft YaHei", "Noto Sans CJK SC"),
        ("宋体", "Noto Serif CJK SC"),
        ("SimSun", "Noto Serif CJK SC"),
        ("Times New Roman", "Liberation Serif"),
    ]
    destination.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(source, "r") as src, zipfile.ZipFile(
        destination, "w", compression=zipfile.ZIP_DEFLATED
    ) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.endswith((".xml", ".rels")):
                text = data.decode("utf-8")
                for old, new in replacements:
                    text = text.replace(old, new)
                data = text.encode("utf-8")
            dst.writestr(item, data)


def main() -> None:
    source = CONTENT.read_text(encoding="utf-8")
    sections = {
        1: section_text(source, "一", "二"),
        2: section_text(source, "二", "三"),
        3: section_text(source, "三", "四"),
        4: section_text(source, "四", "五"),
        5: section_text(source, "五", "六"),
        6: section_text(source, "六", None),
    }
    prose_1, refs_part = sections[1].split("### 参考文献", 1)
    references = []
    for line in refs_part.splitlines():
        match = re.match(r"^\[(\d+)\]\s*(.+)$", line.strip())
        if match:
            references.append(match.group(2).strip())
    if len(references) != 49:
        raise ValueError(f"expected 49 references, found {len(references)}")

    doc = Document(TEMPLATE)
    original_paragraphs = list(doc.paragraphs)

    # Replace stale template metadata with proposal metadata.
    core = doc.core_properties
    core.title = TITLE
    core.subject = "研究生学位论文开题报告"
    core.author = "罗思琪"
    core.last_modified_by = "罗思琪"
    core.comments = ""
    now = datetime.now()
    core.created = now
    core.modified = now
    core.revision = 1

    # Cover table.
    cover = doc.tables[0]
    cover_values = {
        0: "【待填写】",
        1: "罗思琪",
        4: "刘笑宏",
        5: TITLE,
        6: "【待填写】",
        7: "【待填写】",
        8: "【待填写】",
        9: "【待填写】",
    }
    for row_idx, value in cover_values.items():
        insert_run_before_perm_end(cover.cell(row_idx, 1).paragraphs[0], value)
    set_dropdown(cover.cell(2, 1), "学术型博士生 Academic Doctoral Student")
    set_dropdown(cover.cell(3, 1), "全日制 Full-time")

    # Proposal title and source. The source is provisionally set to self-proposed.
    proposal_table = doc.tables[1]
    insert_run_before_perm_end(proposal_table.cell(0, 1).paragraphs[0], TITLE)
    set_source_checkbox(proposal_table.cell(1, 1), selected_index=5)

    # Rebuild official heading runs while preserving their numbered paragraph
    # properties. This also fixes the missing “预期成果” in PDF exports.
    heading_indices = [17, 26, 30, 35, 39, 43]
    for idx in heading_indices:
        rebuild_official_heading(original_paragraphs[idx])
    if "预期成果" not in original_paragraphs[39].text:
        raise ValueError("official heading 5 is missing 预期成果")

    # Keep official numbered headings; add a page break before sections 2–6.
    for idx in heading_indices:
        original_paragraphs[idx].paragraph_format.keep_with_next = True
    for idx in heading_indices[1:]:
        original_paragraphs[idx].paragraph_format.page_break_before = True

    body_slots = {
        1: original_paragraphs[18:22],
        2: original_paragraphs[27:30],
        3: original_paragraphs[31:35],
        4: original_paragraphs[36:39],
        5: original_paragraphs[40:43],
        6: original_paragraphs[44:48],
    }
    fill_range(body_slots[1], parse_blocks(prose_1))
    for sec_no in range(2, 7):
        fill_range(body_slots[sec_no], parse_blocks(sections[sec_no]))

    # Replace reference examples with explicit bracketed numbers and hanging indents.
    ref_slots = [original_paragraphs[23], original_paragraphs[24]]
    ref_clone_source = original_paragraphs[24]
    for idx, reference in enumerate(references):
        if idx < len(ref_slots):
            paragraph = ref_slots[idx]
        else:
            paragraph = clone_blank(ref_clone_source, ref_slots[-1])
            ref_slots.append(paragraph)
        remove_paragraph_numbering(paragraph)
        put_text(paragraph, f"[{idx + 1}] {reference}", "reference")
    original_paragraphs[22].paragraph_format.keep_with_next = True

    # Keep the signature date blank until the student signs the declaration.
    set_signature_date_placeholder(original_paragraphs[50])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    make_render_safe_copy(OUTPUT, RENDER_OUTPUT)
    print(OUTPUT)
    print(RENDER_OUTPUT)


if __name__ == "__main__":
    main()
